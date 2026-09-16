#!/usr/bin/env python3
"""Build original, bounded DOCX fixtures for D431's renderer corpus.

The generated package timestamps may vary, but the authored OOXML content is
deterministic.  This script uses python-docx for ordinary authoring and one
small OOXML patch for real tracked-change markup, which python-docx does not
author directly.
"""

from __future__ import annotations

import struct
import sys
import zlib
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor


OUT = Path(sys.argv[1]).resolve()
BLUE = RGBColor(46, 116, 181)
NAVY = RGBColor(11, 37, 69)


def set_run(run, *, size=11, bold=False, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def base_doc() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    style.font.size = Pt(10.5)
    for name, size in (("Heading 1", 16), ("Heading 2", 13)):
        s = doc.styles[name]
        s.font.name = "Calibri"
        s._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        s._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        s.font.size = Pt(size)
        s.font.color.rgb = BLUE
    return doc


def chrome(section, label: str):
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    header = section.header.paragraphs[0]
    header.text = label
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        set_run(run, size=8, color=RGBColor(100, 100, 100))
    footer = section.footer.paragraphs[0]
    footer.text = "D431 original test fixture — no private content"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in footer.runs:
        set_run(run, size=8, color=RGBColor(100, 100, 100))


def title(doc, text: str, subtitle: str | None = None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    run = p.add_run(text)
    set_run(run, size=24, bold=True, color=NAVY)
    if subtitle:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(16)
        run = p.add_run(subtitle)
        set_run(run, size=12, color=RGBColor(84, 84, 84))


def tiny_original_png(path: Path):
    """Write a small original color-bar PNG without external assets."""
    width, height = 240, 120
    rows = []
    for y in range(height):
        row = bytearray()
        for x in range(width):
            if x < 80:
                rgb = (46, 116, 181)
            elif x < 160:
                rgb = (230, 179, 37)
            else:
                rgb = (11, 37, 69)
            if 35 < y < 85 and 25 < x < 215:
                rgb = tuple(min(255, value + 25) for value in rgb)
            row.extend(rgb)
        rows.append(b"\x00" + bytes(row))
    raw = b"".join(rows)
    def chunk(kind, data):
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data) & 0xffffffff)
    png = b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)) + chunk(b"IDAT", zlib.compress(raw, 9)) + chunk(b"IEND", b"")
    path.write_bytes(png)


def add_tracked_change(paragraph):
    p = paragraph._p
    deleted = OxmlElement("w:del")
    deleted.set(qn("w:id"), "41")
    deleted.set(qn("w:author"), "D431 Fixture Author")
    deleted.set(qn("w:date"), "2026-07-27T00:00:00Z")
    del_run = OxmlElement("w:r")
    del_text = OxmlElement("w:delText")
    del_text.text = "retired wording"
    del_run.append(del_text)
    deleted.append(del_run)
    inserted = OxmlElement("w:ins")
    inserted.set(qn("w:id"), "42")
    inserted.set(qn("w:author"), "D431 Fixture Author")
    inserted.set(qn("w:date"), "2026-07-27T00:00:00Z")
    ins_run = OxmlElement("w:r")
    ins_text = OxmlElement("w:t")
    ins_text.text = "approved wording"
    ins_run.append(ins_text)
    inserted.append(ins_run)
    p.append(deleted)
    p.append(inserted)


def add_external_hyperlink(paragraph, text: str, url: str):
    """Author a real external relationship and w:hyperlink element."""
    rid = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "2E74B5")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.extend([color, underline])
    run.append(props)
    value = OxmlElement("w:t")
    value.text = text
    run.append(value)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def build_simple(path: Path):
    doc = base_doc()
    chrome(doc.sections[0], "D431 / DOCX simple")
    title(doc, "Reader baseline", "One-page original DOCX smoke comparison")
    p = doc.add_paragraph()
    p.add_run("This fixture establishes a compact baseline with ordinary body text.")
    doc.save(path)


def build_rich(path: Path, image_path: Path):
    doc = base_doc()
    chrome(doc.sections[0], "D431 / DOCX rich / portrait")
    title(doc, "OOXML reader qualification brief", "Original content for layout, section, and review-state comparison")
    doc.add_heading("What this fixture proves", level=1)
    p = doc.add_paragraph()
    p.add_run("A resilient reader needs to preserve hierarchy, inline runs, and real document furniture. ")
    add_external_hyperlink(p, "inert D431 fixture link", "https://example.invalid/d431")
    p.add_run(" is a real external OOXML hyperlink; the application must keep it inert or validate it before opening.")
    doc.add_heading("Review checklist", level=2)
    for item in ("Header and footer remain visible.", "Lists retain indentation and wrapping.", "Tables do not collapse their column logic."):
        doc.add_paragraph(item, style="List Bullet")
    table = doc.add_table(rows=1, cols=3)
    table.style = "Light Shading Accent 1"
    headers = ("Surface", "Expected", "Reason")
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
    for row in (("Portrait", "Readable body", "Default reading mode"), ("Landscape", "Wide matrix", "Section geometry"), ("Review", "Tracked change markup", "Revision fidelity")):
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = value
    doc.add_paragraph().add_run("Original embedded color study:").italic = True
    doc.add_picture(str(image_path), width=Inches(4.8))
    note = doc.add_paragraph("Revision state: ")
    add_tracked_change(note)
    landscape = doc.add_section(WD_SECTION.NEW_PAGE)
    landscape.orientation = WD_ORIENT.LANDSCAPE
    landscape.page_width, landscape.page_height = landscape.page_height, landscape.page_width
    landscape.left_margin = Inches(0.65)
    landscape.right_margin = Inches(0.65)
    chrome(landscape, "D431 / DOCX rich / landscape")
    doc.add_heading("Landscape section — comparison matrix", level=1)
    wide = doc.add_table(rows=1, cols=5)
    wide.style = "Light Grid Accent 1"
    for cell, value in zip(wide.rows[0].cells, ("Fixture", "Pages", "Headers", "Tracked", "Notes")):
        cell.text = value
    for i in range(1, 13):
        cells = wide.add_row().cells
        values = (f"Case {i:02d}", str(i % 4 + 1), "yes", "yes" if i % 3 == 0 else "no", "bounded visual evidence")
        for cell, value in zip(cells, values):
            cell.text = value
    doc.add_paragraph("The page intentionally changes orientation so renderers must create a new canvas geometry.")
    portrait = doc.add_section(WD_SECTION.NEW_PAGE)
    portrait.orientation = WD_ORIENT.PORTRAIT
    portrait.page_width, portrait.page_height = Inches(8.5), Inches(11)
    portrait.left_margin = Inches(0.85)
    portrait.right_margin = Inches(0.85)
    chrome(portrait, "D431 / DOCX rich / return to portrait")
    doc.add_heading("Return to portrait", level=1)
    doc.add_paragraph("The final section confirms that the layout can return to the original page geometry after a landscape comparison table.")
    doc.save(path)


def build_edge(path: Path):
    doc = base_doc()
    chrome(doc.sections[0], "D431 / DOCX edge")
    title(doc, "Bounded long-flow fixture", "Stress/edge: 120 short paragraphs and a compact 12-column table")
    doc.add_heading("Long-flow paragraphs", level=1)
    for i in range(1, 121):
        doc.add_paragraph(f"Paragraph {i:03d}: repeated, bounded body content verifies page virtualization without a large compressed source.")
    doc.add_heading("Dense matrix", level=1)
    table = doc.add_table(rows=1, cols=12)
    table.style = "Table Grid"
    for i, cell in enumerate(table.rows[0].cells, 1):
        cell.text = f"C{i}"
    for r in range(12):
        cells = table.add_row().cells
        for c, cell in enumerate(cells, 1):
            cell.text = f"{r + 1}.{c}"
    doc.save(path)


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    image = OUT / "original-color-study.png"
    tiny_original_png(image)
    build_simple(OUT / "simple.docx")
    build_rich(OUT / "rich-mixed-orientation-tracked.docx", image)
    build_edge(OUT / "edge-bounded-flow.docx")
    source = OUT / "rich-mixed-orientation-tracked.docx"
    (OUT / "malformed-truncated.docx").write_bytes(source.read_bytes()[:768])


if __name__ == "__main__":
    main()
